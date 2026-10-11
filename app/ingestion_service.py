"""Validation des sources et extraction bornée, sans exécution de contenu importé."""

import io
import re
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

from app.object_storage import MAX_OBJECT_BYTES

FORMATS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".md": "text/markdown",
    ".html": "text/html",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
MAX_TEXT_CHARS = 100_000
MAX_ZIP_UNCOMPRESSED = 60 * 1024 * 1024


class InvalidSource(ValueError):
    pass


def validate_text(value: str) -> str:
    value = value.strip()
    if not value or len(value) > MAX_TEXT_CHARS or "\x00" in value:
        raise InvalidSource("Contenu texte vide, invalide ou trop volumineux.")
    if sum(ch.isprintable() or ch in "\n\r\t" for ch in value) / len(value) < .98:
        raise InvalidSource("Le contenu contient trop de caractères de contrôle.")
    return value


def _office_members(content: bytes, expected: str) -> None:
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            names = {item.filename for item in infos}
            if len(infos) > 500 or sum(item.file_size for item in infos) > MAX_ZIP_UNCOMPRESSED:
                raise InvalidSource("Archive bureautique trop volumineuse.")
            if "[Content_Types].xml" not in names or expected not in names:
                raise InvalidSource("Contenu bureautique incohérent avec son extension.")
            if any("vbaProject" in name or "externalLinks" in name or name.endswith(".bin")
                   for name in names):
                raise InvalidSource("Contenu bureautique actif ou externe refusé.")
            if any(b'TargetMode="External"' in archive.read(name) or
                   b"TargetMode='External'" in archive.read(name)
                   for name in names if name.endswith(".rels")):
                raise InvalidSource("Relation bureautique externe refusée.")
            if any(item.file_size > 15 * 1024 * 1024 or
                   (item.compress_size and item.file_size / item.compress_size > 100)
                   for item in infos):
                raise InvalidSource("Archive bureautique suspecte.")
    except (BadZipFile, OSError) as exc:
        raise InvalidSource("Archive bureautique invalide.") from exc


def validate_upload(filename: str, content: bytes) -> tuple[str, str]:
    name = PurePath(filename or "").name.strip()
    if not name or len(name) > 180 or name in {".", ".."} or any(ord(c) < 32 for c in name):
        raise InvalidSource("Nom de fichier invalide.")
    suffix = PurePath(name).suffix.lower()
    if suffix not in FORMATS:
        raise InvalidSource("Format de fichier non autorisé.")
    if not content or len(content) > MAX_OBJECT_BYTES:
        raise InvalidSource("Fichier vide ou supérieur à 25 Mio.")
    if suffix == ".pdf" and not content.startswith(b"%PDF-"):
        raise InvalidSource("Signature PDF invalide.")
    if suffix == ".docx":
        _office_members(content, "word/document.xml")
    if suffix == ".xlsx":
        _office_members(content, "xl/workbook.xml")
    if suffix in {".md", ".html"}:
        try:
            validate_text(content.decode("utf-8-sig"))
        except UnicodeDecodeError as exc:
            raise InvalidSource("Le fichier texte doit être encodé en UTF-8.") from exc
    return name, FORMATS[suffix]


def extract_text(filename: str, content: bytes) -> str:
    """Relit les signatures et limite le texte fourni au LLM."""
    _, mime = validate_upload(filename, content)
    if mime == "text/markdown":
        result = content.decode("utf-8-sig")
    elif mime == "text/html":
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(content.decode("utf-8-sig"), "html.parser")
        for tag in soup(["script", "style", "iframe", "object", "svg", "template"]):
            tag.decompose()
        result = soup.get_text("\n", strip=True)
    elif mime == "application/pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content), strict=True)
        if reader.is_encrypted or len(reader.pages) > 100:
            raise InvalidSource("PDF chiffré ou trop long.")
        result = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif mime.endswith("wordprocessingml.document"):
        from docx import Document
        document = Document(io.BytesIO(content))
        result = "\n".join(p.text for p in document.paragraphs)
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            parts = []
            total_chars = 0
            for sheet in workbook.worksheets[:20]:
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= 5000:
                        raise InvalidSource("Tableur trop volumineux.")
                    part = " | ".join(str(cell)[:500] for cell in row if cell is not None)
                    parts.append(part)
                    total_chars += len(part)
                    if total_chars > MAX_TEXT_CHARS:
                        raise InvalidSource("Texte extrait trop volumineux.")
            result = "\n".join(parts)
        finally:
            workbook.close()
    return validate_text(re.sub(r"\r\n?", "\n", result))

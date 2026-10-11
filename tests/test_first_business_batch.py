"""Premier parcours métier : sources, analyse fictive, revue et ACL."""

import asyncio
import io
import json
import os
import subprocess
import sys
import time
from pathlib import Path
from uuid import uuid4

import boto3
import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import select

from app.business_models import AnalysisRun, ContextItem, ContextItemVersion, SourceDocument
from app.identity_models import AuditEvent, ProjectMember
from app.ingestion_service import InvalidSource, extract_text, validate_upload
from app.llm_models import LLMProfileVersion
from app.llm_service import LLMResult, settings as llm_settings
from app.object_storage import ObjectStorage, StorageConfig

from .conftest import create_user, with_db
from .test_identity_acl import csrf, login


class FakeS3:
    def __init__(self):
        self.objects = {}

    def put_object(self, **kwargs):
        self.objects[kwargs["Key"]] = kwargs["Body"]

    def get_object(self, **kwargs):
        return {"Body": io.BytesIO(self.objects[kwargs["Key"]])}

    def delete_object(self, **kwargs):
        self.objects.pop(kwargs["Key"], None)


def test_validation_des_formats_et_extraction_textuelle():
    name, mime = validate_upload("besoin.md", b"# Besoin\nUn parcours utilisateur.")
    assert name == "besoin.md" and mime == "text/markdown"
    assert "parcours" in extract_text(name, b"# Besoin\nUn parcours utilisateur.")
    assert extract_text("page.html", b"<script>secret</script><h1>Besoin</h1>") == "Besoin"
    document = Document()
    document.add_paragraph("Besoin DOCX")
    docx = io.BytesIO()
    document.save(docx)
    assert "Besoin DOCX" in extract_text("besoin.docx", docx.getvalue())
    workbook = Workbook()
    workbook.active["A1"] = "Besoin XLSX"
    xlsx = io.BytesIO()
    workbook.save(xlsx)
    assert "Besoin XLSX" in extract_text("besoin.xlsx", xlsx.getvalue())
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = writer._add_object(DictionaryObject({NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")}))
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
        DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (Besoin PDF) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    pdf = io.BytesIO()
    writer.write(pdf)
    assert "Besoin PDF" in extract_text("besoin.pdf", pdf.getvalue())
    for filename, content in [("fichier.exe", b"x"), ("fichier.pdf", b"texte"),
                              ("fichier.docx", b"PK faux"), ("fichier.xlsx", b"PK faux")]:
        try:
            validate_upload(filename, content)
        except InvalidSource:
            pass
        else:
            raise AssertionError(f"Format invalide accepté : {filename}")


def test_import_analyse_revue_et_permissions(pg_client, monkeypatch):
    monkeypatch.setenv("MIMIC_REDIS_URL", "redis://127.0.0.1:56379/15")
    monkeypatch.setenv("MIMIC_LLM_ALLOWED_BASE_URLS", "http://127.0.0.1:11434")
    from app import worker

    fake = FakeS3()
    storage = ObjectStorage(StorageConfig("https://s3.example.org", "mimic-test", "us-east-1", "a", "b"), fake)
    monkeypatch.setattr(ObjectStorage, "from_env", classmethod(lambda cls: storage))
    monkeypatch.setattr(worker.extract_source, "delay", lambda source_id: None)
    monkeypatch.setattr(worker.analyze_source, "delay", lambda run_id: None)

    admin_id, admin_email = create_user(admin=True)
    token = login(pg_client, admin_email)
    response = pg_client.post("/projects", data={"name": "Premier lot", "description": "",
                                                  "csrf_token": token})
    project_id = response.headers["location"].split("/")[-1]
    assert pg_client.get(f"/projects/{project_id}/sources").status_code == 200
    assert pg_client.post(f"/projects/{project_id}/sources", data={"csrf_token": "faux"},
                          files={"file": ("besoin.md", b"# Besoin")}).status_code == 403
    assert pg_client.post(f"/projects/{project_id}/sources", data={"csrf_token": token},
                          files={"file": ("malware.exe", b"x")}).status_code == 422
    uploaded = pg_client.post(f"/projects/{project_id}/sources", data={"csrf_token": token},
                              files={"file": ("besoin.md", b"# Besoin\nUn acteur utilise le produit.")})
    assert uploaded.status_code == 303
    source_id = uploaded.headers["location"].split("/")[-1]
    assert pg_client.get(f"/projects/{project_id}/sources/{source_id}/download").content.startswith(b"# Besoin")
    asyncio.run(worker._extract_source(source_id))
    assert pg_client.get(f"/projects/{project_id}/sources/{source_id}").status_code == 200

    async def profile(db):
        row = LLMProfileVersion(profile_id=str(uuid4()), version=1, name="Local", provider="ollama",
                                model_name="test", base_url="http://127.0.0.1:11434", api_version=None,
                                credential_env=None, temperature=None, max_output_tokens=1000,
                                timeout_seconds=10, max_retries=0, enabled=True, is_current=True,
                                created_by=admin_id)
        db.add(row)
        await db.flush()
        settings = await llm_settings(db)
        settings.default_version_id = row.id
        return row.id

    profile_id = asyncio.run(with_db(profile))

    async def fake_invoke(profile, messages, **kwargs):
        assert profile.id == profile_id
        assert "acteur" in messages[1]["content"]
        return LLMResult(json.dumps({"items": [{"type": "actor", "title": "Utilisateur",
            "content": "Utilise le produit."}]}), profile.id, "ollama", "test", 30, 20)

    monkeypatch.setattr(worker, "invoke_profile", fake_invoke)
    started = pg_client.post(f"/projects/{project_id}/sources/{source_id}/analyze",
                             data={"csrf_token": token})
    assert started.status_code == 303
    duplicate = pg_client.post(f"/projects/{project_id}/sources/{source_id}/analyze",
                               data={"csrf_token": token})
    assert duplicate.status_code == 409

    async def state(db):
        source = await db.get(SourceDocument, source_id)
        run = await db.scalar(select(AnalysisRun).where(AnalysisRun.source_id == source_id))
        return source.status, run.id, run.llm_profile_version_id

    source_status, run_id, pinned = asyncio.run(with_db(state))
    assert source_status == "analyzing" and pinned == profile_id
    asyncio.run(worker._analyze_source(run_id))

    async def proposal(db):
        item = await db.scalar(select(ContextItem).where(ContextItem.project_id == project_id))
        run = await db.get(AnalysisRun, run_id)
        return item.id, item.status, item.source_id, run.status, run.input_tokens

    item_id, status, linked_source, run_status, tokens = asyncio.run(with_db(proposal))
    assert (status, linked_source, run_status, tokens) == ("Brouillon IA", source_id, "completed", 30)
    assert pg_client.get(f"/projects/{project_id}/context?item_id={item_id}").status_code == 200

    reader_id, reader_email = create_user()
    async def add_reader(db):
        db.add(ProjectMember(project_id=project_id, user_id=reader_id, role_code="reader", is_active=True))
    asyncio.run(with_db(add_reader))
    admin_session = pg_client.cookies.get("mimic-session")
    pg_client.cookies.clear()
    reader_token = login(pg_client, reader_email)
    assert pg_client.get(f"/projects/{project_id}/sources/{source_id}").status_code == 200
    assert pg_client.post(f"/projects/{project_id}/context/{item_id}/decisions", data={
        "decision": "review", "expected_version": 1, "csrf_token": reader_token}).status_code == 403
    assert pg_client.post(f"/projects/{project_id}/sources/{source_id}/analyze", data={
        "csrf_token": reader_token}).status_code == 403

    pg_client.cookies.set("mimic-session", admin_session)
    assert pg_client.post(f"/projects/{project_id}/context/{item_id}/decisions", data={
        "decision": "review", "expected_version": 1, "csrf_token": token}).status_code == 303
    assert pg_client.post(f"/projects/{project_id}/context/{item_id}/decisions", data={
        "decision": "validate", "expected_version": 1, "csrf_token": token}).status_code == 409
    assert pg_client.post(f"/projects/{project_id}/context/{item_id}/decisions", data={
        "decision": "validate", "expected_version": 2, "csrf_token": token}).status_code == 303

    async def final_state(db):
        item = await db.get(ContextItem, item_id)
        versions = (await db.scalars(select(ContextItemVersion).where(ContextItemVersion.item_id == item_id)
                                     .order_by(ContextItemVersion.version))).all()
        actions = (await db.scalars(select(AuditEvent.action).where(AuditEvent.project_id == project_id))).all()
        return item, versions, actions

    item, versions, actions = asyncio.run(with_db(final_state))
    assert item.status == "Validé" and [v.status for v in versions] == ["Brouillon IA", "À revoir", "Validé"]
    assert {"source.upload", "source.extract", "source.analyze", "context.review", "context.validate"} <= set(actions)

    async def bad_invoke(profile, messages, **kwargs):
        return LLMResult("réponse invalide", profile.id, "ollama", "test", None, None)

    monkeypatch.setattr(worker, "invoke_profile", bad_invoke)
    assert pg_client.post(f"/projects/{project_id}/sources/{source_id}/analyze",
                          data={"csrf_token": token}).status_code == 303
    async def new_run(db):
        return await db.scalar(select(AnalysisRun.id).where(AnalysisRun.source_id == source_id)
                               .order_by(AnalysisRun.created_at.desc()).limit(1))
    failed_run_id = asyncio.run(with_db(new_run))
    asyncio.run(worker._analyze_source(failed_run_id))
    async def failed_state(db):
        run = await db.get(AnalysisRun, failed_run_id)
        source = await db.get(SourceDocument, source_id)
        count = len((await db.scalars(select(ContextItem).where(ContextItem.source_id == source_id))).all())
        return run.status, run.error_code, source.status, count
    assert asyncio.run(with_db(failed_state)) == ("failed", "invalid_response", "extracted", 1)

    technical = pg_client.post(f"/projects/{project_id}/context", data={
        "type": "component", "title": "API", "content": "Composant applicatif.", "csrf_token": token})
    business = pg_client.post(f"/projects/{project_id}/context", data={
        "type": "requirement", "title": "Exigence", "content": "Besoin métier.", "csrf_token": token})
    technical_id = technical.headers["location"].split("item_id=")[-1]
    business_id = business.headers["location"].split("item_id=")[-1]
    po_id, po_email = create_user()
    architect_id, architect_email = create_user()
    async def add_reviewers(db):
        db.add(ProjectMember(project_id=project_id, user_id=po_id, role_code="product_owner", is_active=True))
        db.add(ProjectMember(project_id=project_id, user_id=architect_id, role_code="architect", is_active=True))
    asyncio.run(with_db(add_reviewers))
    pg_client.cookies.clear()
    po_token = login(pg_client, po_email)
    assert pg_client.post(f"/projects/{project_id}/context/{technical_id}/decisions", data={
        "decision": "validate", "expected_version": 1, "csrf_token": po_token}).status_code == 403
    assert pg_client.post(f"/projects/{project_id}/context/{business_id}/decisions", data={
        "decision": "validate", "expected_version": 1, "csrf_token": po_token}).status_code == 303
    pg_client.cookies.clear()
    architect_token = login(pg_client, architect_email)
    assert pg_client.post(f"/projects/{project_id}/context/{technical_id}/decisions", data={
        "decision": "validate", "expected_version": 1, "csrf_token": architect_token}).status_code == 303
    assert pg_client.post(f"/projects/{project_id}/context/{business_id}/decisions", data={
        "decision": "validate", "expected_version": 2, "csrf_token": architect_token}).status_code == 403


def test_import_reel_via_celery_et_minio(pg_client, monkeypatch):
    endpoint = os.environ.get("MIMIC_TEST_S3_ENDPOINT_URL")
    redis_url = os.environ.get("MIMIC_TEST_REDIS_URL")
    access_key = os.environ.get("MIMIC_S3_ACCESS_KEY")
    secret_key = os.environ.get("MIMIC_S3_SECRET_KEY")
    if not all((endpoint, redis_url, access_key, secret_key)):
        pytest.skip("Pile locale MinIO/Redis non configurée")
    bucket = f"mimic-business-{uuid4().hex[:12]}"
    monkeypatch.setenv("MIMIC_S3_ENDPOINT_URL", endpoint)
    monkeypatch.setenv("MIMIC_S3_BUCKET", bucket)
    monkeypatch.setenv("MIMIC_S3_ALLOW_HTTP", "true")
    monkeypatch.setenv("MIMIC_REDIS_URL", redis_url)
    s3 = boto3.client("s3", endpoint_url=endpoint, region_name="us-east-1",
                      aws_access_key_id=access_key, aws_secret_access_key=secret_key)
    s3.create_bucket(Bucket=bucket)
    from app.worker import celery_app
    celery_app.conf.broker_url = redis_url
    celery_app.conf.result_backend = redis_url
    command = [sys.executable, "-m", "celery", "-A", "app.worker:celery_app", "worker",
               "--pool=solo", "--concurrency=1", "--loglevel=WARNING", "--hostname=mimic-business@%h"]
    worker = subprocess.Popen(command, cwd=Path(__file__).resolve().parents[1], env=os.environ.copy(),
                              stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    key = None
    try:
        for _ in range(20):
            if celery_app.control.ping(timeout=1):
                break
            time.sleep(.25)
        else:
            pytest.fail("Worker Celery indisponible")
        _, email = create_user(admin=True)
        token = login(pg_client, email)
        project = pg_client.post("/projects", data={"name": "Test worker", "csrf_token": token})
        project_id = project.headers["location"].split("/")[-1]
        response = pg_client.post(f"/projects/{project_id}/sources",
            data={"csrf_token": token}, files={"file": ("besoin.md", b"# Source\nExtraction asynchrone.")})
        assert response.status_code == 303
        source_id = response.headers["location"].split("/")[-1]

        async def state(db):
            source = await db.get(SourceDocument, source_id)
            return source.status, source.extracted_text, source.object_key

        for _ in range(40):
            status, extracted, key = asyncio.run(with_db(state))
            if status == "extracted":
                break
            time.sleep(.5)
        assert status == "extracted" and "Extraction asynchrone" in extracted
    finally:
        worker.terminate()
        try:
            worker.wait(timeout=10)
        except subprocess.TimeoutExpired:
            worker.kill()
            worker.wait(timeout=5)
        if key:
            s3.delete_object(Bucket=bucket, Key=key)
        s3.delete_bucket(Bucket=bucket)
